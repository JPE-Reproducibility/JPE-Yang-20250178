from __future__ import annotations

import os
import sys
import warnings
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
MPLCONFIGDIR = ROOT / "output" / "logs" / "matplotlib-cache"
MPLCONFIGDIR.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(MPLCONFIGDIR))

import matplotlib
import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from matplotlib.ticker import PercentFormatter

matplotlib.use("Agg")
import matplotlib.pyplot as plt

sys.path.insert(0, str(ROOT / "code" / "auxiliary"))
import paths

warnings.filterwarnings("ignore")

IPO_FILE = paths.MANUAL_DIR_CONF / "figure1c_IPO Share Table.xlsx"
IPO_FALLBACK_FILE = paths.MANUAL_DIR_CONF / "figure1c_IPO Share Table 1c.xlsx"
IPO_CORRECTED_FILES = {
    "em": paths.MANUAL_DIR_CONF / "figure1c_IPO_CapitalIQ_to_PB_EM_0901,Final.xlsx",
    "us": paths.MANUAL_DIR_CONF / "figure1c_IPO_CapitalIQ_to_PB_US_Final.xlsx",
    "dm_ex_us": paths.MANUAL_DIR_CONF
    / "figure1c_IPO_CapitalIQ_to_PB_nonEM_nonUS.With JL Additions.xlsx",
}
IPO_GENERATED_TABLE = paths.INTERMEDIATE_DIR / "figure1c_IPO Share Table.generated.xlsx"
IPO_TABULATIONS_FILE = paths.INTERMEDIATE_DIR / "figure1c_ipo_tabulations.csv"
IPO_PANEL_FILE = paths.INTERMEDIATE_DIR / "figure1c_generated_ipo_shares.csv"
IPO_COMPARISON_FILE = paths.INTERMEDIATE_DIR / "figure1c_generated_vs_excel_chart_block.csv"
FIG1AB_FILE = paths.MANUAL_DIR_CONF / "figure1ab_Figures 1a and 1b.xlsx"

REGIONS = ["US", "Developed ex US", "China", "Developing ex China"]
REGION_COLORS = {
    "US": "#4E71BE",
    "Developed ex US": "#DE8344",
    "China": "#EA3323",
    "Developing ex China": "#F6C242",
}
YEARS = list(range(2001, 2022))
GRID_COLOR = "#E6E6E6"
AXIS_COLOR = "#595959"
CHINA_HEADQUARTERS = {"China", "Hong Kong", "Macau"}
IPO_GROUPS = ["China", "All Other EM", "U.S.", "All Other DM"]
IPO_METRICS = ["IPO Count", "Market Capitalization", "R&D Expenditures"]
NON_ENT_INDUSTRIES = {
    "Asset Management and Custody Banks",
    "Coal and Consumable Fuels",
    "Commercial and Residential Mortgage Finance",
    "Copper",
    "Data Center REITs",
    "Diversified Banks",
    "Diversified Capital Markets",
    "Diversified Financial Services",
    "Diversified Metals and Mining",
    "Diversified REITs",
    "Electric Utilities",
    "Gas Utilities",
    "Gold",
    "Health Care REITs",
    "Hotel and Resort REITs",
    "Independent Power Producers and Energy Traders",
    "Industrial REITs",
    "Insurance Brokers",
    "Integrated Oil and Gas",
    "Investment Banking and Brokerage",
    "Life and Health Insurance",
    "Mortgage REITs",
    "Multi-Family Residential REITs",
    "Multi-Utilities",
    "Multi-line Insurance",
    "Office REITs",
    "Oil and Gas Drilling",
    "Oil and Gas Exploration and Production",
    "Oil and Gas Refining and Marketing",
    "Oil and Gas Storage and Transportation",
    "Other Specialized REITs",
    "Precious Metals and Minerals",
    "Property and Casualty Insurance",
    "Regional Banks",
    "Reinsurance",
    "Retail REITs",
    "Self-Storage REITs",
    "Specialized Finance",
    "Steel",
    "Water Utilities",
    "Oil and Gas Equipment and Services",
    "Single-Family Residential REITs",
    "Telecom Tower REITs",
}

plt.rcParams.update(
    {
        "font.family": ["Arial", "DejaVu Sans"],
        "axes.edgecolor": "black",
        "axes.labelcolor": AXIS_COLOR,
        "xtick.color": "black",
        "ytick.color": "black",
        "legend.frameon": False,
        "pdf.fonttype": 42,
        "ps.fonttype": 42,
    }
)

def _require(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"Required input not found: {path}")
    if path.stat().st_size == 0:
        raise ValueError(f"Required input is empty: {path}")

def _num(value):
    try:
        if value in ("", None):
            return None
        return float(value)
    except (TypeError, ValueError):
        return None

def _style_excel_axes(ax, include_grid: bool = True) -> None:
    if include_grid:
        ax.grid(axis="y", color=GRID_COLOR, linewidth=0.8)
        ax.set_axisbelow(True)
    for spine in ax.spines.values():
        spine.set_color("black")
        spine.set_linewidth(0.8)

def _resolve_column(columns: list[str], stata_prefix: str) -> str:
    matches = [col for col in columns if str(col).startswith(stata_prefix)]
    if len(matches) != 1:
        raise ValueError(f"Expected one column beginning with {stata_prefix!r}; found {matches}.")
    return matches[0]

def load_vc_levels() -> pd.DataFrame:
    _require(FIG1AB_FILE)
    wb = openpyxl.load_workbook(FIG1AB_FILE, data_only=True, read_only=True)

    pitchbook = {}
    ws = wb["Original PitchBook Data"]
    for row in ws.iter_rows(values_only=True):
        year = _num(row[1] if len(row) > 1 else None)
        if year is None or int(year) != year:
            continue
        year = int(year)
        global_vc = _num(row[3] if len(row) > 3 else None)
        developed_ex_us = _num(row[4] if len(row) > 4 else None)
        us = _num(row[5] if len(row) > 5 else None)
        china = _num(row[6] if len(row) > 6 else None)
        if None in (global_vc, developed_ex_us, us, china):
            continue
        pitchbook[year] = {
            "US": us,
            "Developed ex US": developed_ex_us,
            "China": china,
            "Developing ex China": global_vc - developed_ex_us - us - china,
        }

    bea = wb["BEA Stats"]
    year_header = list(next(bea.iter_rows(min_row=6, max_row=6, values_only=True)))
    gdp_deflator_row = list(next(bea.iter_rows(min_row=8, max_row=8, values_only=True)))
    deflators = {}
    for year, value in zip(year_header, gdp_deflator_row):
        year_num = _num(year)
        value_num = _num(value)
        if year_num is not None and value_num is not None and int(year_num) == year_num:
            deflators[int(year_num)] = value_num
    wb.close()

    if 2011 not in deflators:
        raise ValueError("BEA GDP deflator row does not contain 2011.")

    rows = []
    for year in YEARS:
        if year not in pitchbook or year not in deflators:
            raise ValueError(f"Missing PitchBook or BEA source value for {year}.")
        price_index = deflators[year] / deflators[2011]
        row = {"year": year}
        for region in REGIONS:
            row[region] = pitchbook[year][region] / price_index / 1000.0
        rows.append(row)

    df = pd.DataFrame(rows).set_index("year")
    df.to_csv(paths.INTERMEDIATE_DIR / "figure1ab_real_vc_2011usd.csv")
    return df[REGIONS]

def figure1a(levels: pd.DataFrame) -> None:
    shares = levels.div(levels.sum(axis=1), axis=0)
    fig, ax = plt.subplots(figsize=(10, 4.5))
    ax.stackplot(
        shares.index,
        [shares[r] for r in REGIONS],
        labels=REGIONS,
        colors=[REGION_COLORS[r] for r in REGIONS],
        linewidth=0,
    )
    ax.set_xlim(shares.index.min(), shares.index.max())
    ax.set_ylim(0, 1)
    ax.set_yticks([i / 10 for i in range(0, 11)])
    ax.yaxis.set_major_formatter(PercentFormatter(xmax=1, decimals=0))
    ax.set_xticks(YEARS)
    ax.set_xticklabels(YEARS, rotation=0, fontsize=8)
    ax.margins(x=0)
    _style_excel_axes(ax, include_grid=False)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=4, frameon=False)
    fig.tight_layout()
    out = paths.FIGURE_DIR / "figure1a.pdf"
    fig.savefig(out, dpi=150)
    plt.close(fig)
    print(f"Saved {out}")

def figure1b(levels: pd.DataFrame) -> None:
    fig, ax = plt.subplots(figsize=(10, 4.5))
    for r in REGIONS:
        ax.plot(levels.index, levels[r], color=REGION_COLORS[r], linewidth=1.8, label=r)
    ax.set_xlim(levels.index.min(), levels.index.max())
    ax.set_xticks(YEARS)
    ax.set_xticklabels(YEARS, rotation=0, fontsize=8)
    ax.set_ylim(0, 300)
    ax.set_yticks(range(0, 301, 50))
    ax.set_title("Billions USD", loc="left", fontsize=11)
    _style_excel_axes(ax)
    ax.legend(loc="upper center", ncol=4, frameon=False)
    fig.tight_layout()
    out = paths.FIGURE_DIR / "figure1b.pdf"
    fig.savefig(out, dpi=150)
    plt.close(fig)
    print(f"Saved {out}")

def _read_corrected_ipo_workbook(path: Path, source_sample: str) -> pd.DataFrame:
    _require(path)
    raw = pd.read_excel(path, sheet_name=0, engine="openpyxl")
    columns = list(raw.columns)
    pb_vc_col = _resolve_column(columns, "PB_VC")
    headquarters_col = _resolve_column(columns, "Headquarters")
    market_col = _resolve_column(columns, "Market")
    rd_col = _resolve_column(columns, "RDExpenseFY2022")

    df = pd.DataFrame(
        {
            "source_sample": source_sample,
            "source_file": path.name,
            "PB_VC": pd.to_numeric(raw[pb_vc_col], errors="coerce"),
            "Headquarters": raw[headquarters_col].astype("string").str.strip(),
            "CompanyIndustry": raw["CompanyIndustry"].astype("string").str.strip(),
            "Market": pd.to_numeric(raw[market_col], errors="coerce"),
            "RD": pd.to_numeric(raw[rd_col], errors="coerce"),
        }
    )
    df["NonEnt"] = df["CompanyIndustry"].isin(NON_ENT_INDUSTRIES)
    return df

def _safe_share(numerator: float, denominator: float) -> float:
    return float("nan") if denominator == 0 else numerator / denominator

def _summarize_ipo_group(group: str, source_sample: str, sample: pd.DataFrame) -> dict:
    sample = sample[sample["PB_VC"].isin([0, 1])].copy()
    non_vc = sample["PB_VC"] == 0
    vc = sample["PB_VC"] == 1
    count_non_vc = int(non_vc.sum())
    count_vc = int(vc.sum())
    market_non_vc = float(sample.loc[non_vc, "Market"].sum(skipna=True))
    market_vc = float(sample.loc[vc, "Market"].sum(skipna=True))
    rd_non_vc = float(sample.loc[non_vc, "RD"].sum(skipna=True))
    rd_vc = float(sample.loc[vc, "RD"].sum(skipna=True))

    return {
        "Region": group,
        "Source sample": source_sample,
        "Total IPO count": count_non_vc + count_vc,
        "Non-VC IPO count": count_non_vc,
        "VC IPO count": count_vc,
        "IPO Count": _safe_share(count_vc, count_non_vc + count_vc),
        "Non-VC market capitalization": market_non_vc,
        "VC market capitalization": market_vc,
        "Market Capitalization": _safe_share(market_vc, market_non_vc + market_vc),
        "Non-VC R&D expenditures": rd_non_vc,
        "VC R&D expenditures": rd_vc,
        "R&D Expenditures": _safe_share(rd_vc, rd_non_vc + rd_vc),
    }

def _read_ipo_chart_block(source_file: Path) -> pd.DataFrame:
    _require(source_file)
    wb = openpyxl.load_workbook(source_file, data_only=True, read_only=True)
    ws = wb["IPO Share Table"] if "IPO Share Table" in wb.sheetnames else wb["Sheet1"]
    metrics = [ws["N2"].value, ws["O2"].value, ws["P2"].value]
    values = {}
    for row in range(3, 7):
        group = ws[f"M{row}"].value
        if not group:
            continue
        values[group] = {
            metric: _num(ws.cell(row=row, column=col).value)
            for metric, col in zip(metrics, range(14, 17))
        }
    wb.close()

    panel = pd.DataFrame.from_dict(values, orient="index").reindex(IPO_GROUPS)[IPO_METRICS]
    if panel.isna().any().any():
        raise ValueError(f"Missing Figure 1c chart-block values in {source_file}.")
    return panel

def _write_generated_ipo_share_table(summary: pd.DataFrame, panel: pd.DataFrame) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "IPO Share Table"
    ws["A1"] = "Generated Figure 1C IPO share table"
    ws["A2"] = (
        "Inputs are the three manually corrected Capital IQ/PitchBook IPO workbooks."
    )
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"].alignment = Alignment(wrap_text=True)

    headers = [
        "Region",
        "Source sample",
        "Total IPO count",
        "Non-VC IPO count",
        "VC IPO count",
        "VC share of IPO count",
        "Non-VC market capitalization",
        "VC market capitalization",
        "VC share of market capitalization",
        "Non-VC R&D expenditures",
        "VC R&D expenditures",
        "VC share of R&D expenditures",
    ]
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")

    for row_idx, record in enumerate(summary.to_dict("records"), start=5):
        values = [
            record["Region"],
            record["Source sample"],
            record["Total IPO count"],
            record["Non-VC IPO count"],
            record["VC IPO count"],
            record["IPO Count"],
            record["Non-VC market capitalization"],
            record["VC market capitalization"],
            record["Market Capitalization"],
            record["Non-VC R&D expenditures"],
            record["VC R&D expenditures"],
            record["R&D Expenditures"],
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_idx, column=col, value=value)

    for col in [6, 9, 12, 14, 15, 16]:
        for row in range(3, 9):
            ws.cell(row=row, column=col).number_format = "0.0%"
    for col in [3, 4, 5]:
        for row in range(5, 9):
            ws.cell(row=row, column=col).number_format = "#,##0"
    for col in [7, 8, 10, 11]:
        for row in range(5, 9):
            ws.cell(row=row, column=col).number_format = "#,##0.0"

    ws["M2"] = ""
    for col, metric in enumerate(IPO_METRICS, start=14):
        cell = ws.cell(row=2, column=col, value=metric)
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for row_idx, group in enumerate(IPO_GROUPS, start=3):
        ws.cell(row=row_idx, column=13, value=group)
        for col_idx, metric in enumerate(IPO_METRICS, start=14):
            ws.cell(row=row_idx, column=col_idx, value=float(panel.loc[group, metric]))
            ws.cell(row=row_idx, column=col_idx).number_format = "0.0%"

    for col in range(1, 17):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 18
    ws.freeze_panes = "A5"

    inputs = wb.create_sheet("Inputs")
    inputs.append(["Role", "Package file", "Note"])
    inputs.append(
        [
            "Manual corrected EM input",
            IPO_CORRECTED_FILES["em"].name,
            "Includes manual Capital IQ/PitchBook corrections; split into China and all other EM.",
        ]
    )
    inputs.append(
        [
            "Manual corrected U.S. input",
            IPO_CORRECTED_FILES["us"].name,
            "Includes manual Capital IQ/PitchBook corrections.",
        ]
    )
    inputs.append(
        [
            "Manual corrected all other DM input",
            IPO_CORRECTED_FILES["dm_ex_us"].name,
            "Includes Josh Lerner additions/manual research corrections.",
        ]
    )
    inputs.append(
        [
            "Reference Excel chart workbook",
            IPO_FILE.name,
            "Used only to cross-check the generated values.",
        ]
    )
    for cell in inputs[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAF7")
    for col in range(1, 4):
        inputs.column_dimensions[get_column_letter(col)].width = 36

    wb.save(IPO_GENERATED_TABLE)

def _write_ipo_reference_comparison(panel: pd.DataFrame) -> None:
    reference_file = IPO_FILE if IPO_FILE.exists() else IPO_FALLBACK_FILE
    if not reference_file.exists():
        return
    reference = _read_ipo_chart_block(reference_file)
    rows = []
    for group in IPO_GROUPS:
        for metric in IPO_METRICS:
            rows.append(
                {
                    "Group": group,
                    "Metric": metric,
                    "Generated": panel.loc[group, metric],
                    "Reference Excel chart block": reference.loc[group, metric],
                    "Difference": panel.loc[group, metric] - reference.loc[group, metric],
                }
            )
    pd.DataFrame(rows).to_csv(IPO_COMPARISON_FILE, index=False)

def build_ipo_share_table() -> pd.DataFrame:
    em = _read_corrected_ipo_workbook(IPO_CORRECTED_FILES["em"], "Corrected EM")
    us = _read_corrected_ipo_workbook(IPO_CORRECTED_FILES["us"], "Corrected U.S.")
    dm_ex_us = _read_corrected_ipo_workbook(IPO_CORRECTED_FILES["dm_ex_us"], "Corrected all other DM")

    em_kept = em.loc[~em["NonEnt"]].copy()
    us_kept = us.loc[~us["NonEnt"]].copy()
    dm_kept = dm_ex_us.loc[~dm_ex_us["NonEnt"]].copy()

    summaries = [
        _summarize_ipo_group(
            "All Other EM",
            "Corrected EM, excluding China/Hong Kong/Macau",
            em_kept.loc[~em_kept["Headquarters"].isin(CHINA_HEADQUARTERS)],
        ),
        _summarize_ipo_group(
            "China",
            "Corrected EM, China/Hong Kong/Macau only",
            em_kept.loc[em_kept["Headquarters"].isin(CHINA_HEADQUARTERS)],
        ),
        _summarize_ipo_group("U.S.", "Corrected U.S.", us_kept),
        _summarize_ipo_group("All Other DM", "Corrected all other DM", dm_kept),
    ]
    summary = pd.DataFrame(summaries)
    summary.to_csv(IPO_TABULATIONS_FILE, index=False)

    panel = summary.set_index("Region").reindex(IPO_GROUPS)[IPO_METRICS]
    _write_generated_ipo_share_table(summary, panel)
    _write_ipo_reference_comparison(panel)
    return panel

def load_ipo_panel() -> pd.DataFrame:
    build_ipo_share_table()
    panel = _read_ipo_chart_block(IPO_GENERATED_TABLE)
    panel.to_csv(IPO_PANEL_FILE)
    return panel

def figure1c() -> None:
    block = load_ipo_panel()
    metrics = ["IPO Count", "Market Capitalization", "R&D Expenditures"]
    groups = ["China", "All Other EM", "U.S.", "All Other DM"]

    group_colors = {
        "China": "#FF0000",
        "All Other EM": "#92D050",
        "U.S.": "#0070C0",
        "All Other DM": "#FFC000",
    }
    hatches = {"China": "", "All Other EM": "xx", "U.S.": "", "All Other DM": "---"}

    import numpy as np

    n_groups = len(groups)
    x = np.arange(len(metrics))
    width = 0.17
    fig, ax = plt.subplots(figsize=(11, 5))
    for i, g in enumerate(groups):
        offset = (i - (n_groups - 1) / 2) * width
        facecolor = group_colors[g] if hatches[g] == "" else "white"
        edgecolor = group_colors[g]
        ax.bar(
            x + offset,
            block.loc[g, metrics].values,
            width,
            label=g,
            color=facecolor,
            hatch=hatches[g],
            edgecolor=edgecolor,
            linewidth=0.8,
        )
    ax.set_title("VC-Backed Firms as a Share of Young Public Firms", fontsize=12, fontweight="bold")
    ax.set_xticks(x)
    ax.set_xticklabels(metrics, fontsize=12)
    ax.set_ylim(0, 0.9)
    ax.set_yticks([i / 10 for i in range(0, 10)])
    ax.yaxis.set_major_formatter(PercentFormatter(xmax=1, decimals=0))
    _style_excel_axes(ax)
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.08), ncol=4, frameon=False)
    fig.tight_layout()
    out = paths.FIGURE_DIR / "figure1c.pdf"
    fig.savefig(out, dpi=150)
    plt.close(fig)
    print(f"Saved {out}")

def main() -> None:
    levels = load_vc_levels()
    figure1a(levels)
    figure1b(levels)
    figure1c()

if __name__ == "__main__":
    main()
