from __future__ import annotations

import os
from pathlib import Path

import matplotlib
import matplotlib.pyplot as plt
import openpyxl
import pandas as pd

matplotlib.use("Agg")

os.chdir(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "confidential-data-not-for-publication" / "Raw" / "other_data_resource"
OUT = ROOT / "output" / "figures" / "figureA4.pdf"
INTERMEDIATE = ROOT / "output" / "intermediate" / "figureA4_china_vc_crossvalidation.csv"
OUT.parent.mkdir(parents=True, exist_ok=True)
INTERMEDIATE.parent.mkdir(parents=True, exist_ok=True)

COMPARISON_FILE = RAW / "figureA4_Comparisons of PB and Other Sources.xlsx"
ZERO2IPO_FILE = RAW / "figureA4_Data from Zero2IPO from HBS China Research Center.xlsx"

def _num(value):
    try:
        if value in ("", None):
            return None
        return float(value)
    except (TypeError, ValueError):
        return None

def _require(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"Required input not found: {path}")
    if path.stat().st_size == 0:
        raise ValueError(f"Required input is empty: {path}")

def load_pb_cvi() -> pd.DataFrame:
    _require(COMPARISON_FILE)
    wb = openpyxl.load_workbook(COMPARISON_FILE, data_only=True, read_only=True)
    ws = wb["Data"]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    records = []
    cvi_count_share = None
    cvi_size_share = None
    for row in rows[1:]:
        year = _num(row[0] if len(row) > 0 else None)
        if year is None or int(year) != year:
            continue
        year = int(year)
        if year == 2018:
            cvi_count_share = _num(row[5] if len(row) > 5 else None)
            cvi_size_share = _num(row[8] if len(row) > 8 else None)
        records.append(row)

    if cvi_count_share is None or cvi_size_share is None:
        raise ValueError("Could not find 2018 CVI shares in comparison workbook.")

    out = []
    for row in records:
        year = int(_num(row[0]))
        pb_size_m = _num(row[2] if len(row) > 2 else None)
        cvi_vcpe_size = _num(row[7] if len(row) > 7 else None)
        cvi_share = cvi_size_share
        out.append(
            {
                "Year": year,
                "PitchBook": pb_size_m / 1000.0 if pb_size_m is not None else None,
                "CVSource / ChinaVenture": (
                    cvi_vcpe_size * cvi_share
                    if cvi_vcpe_size is not None and cvi_share is not None
                    else None
                ),
            }
        )
    return pd.DataFrame(out)

def load_zero2ipo() -> pd.DataFrame:
    _require(ZERO2IPO_FILE)
    wb = openpyxl.load_workbook(ZERO2IPO_FILE, data_only=True, read_only=True)
    ws = wb["Data"]
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        year = _num(row[0] if len(row) > 0 else None)
        size_rmb = _num(row[2] if len(row) > 2 else None)
        rmb_per_usd = _num(row[4] if len(row) > 4 else None)
        if year is None or int(year) != year:
            continue
        records.append(
            {
                "Year": int(year),
                "Zero2IPO": (
                    size_rmb / rmb_per_usd
                    if size_rmb is not None and rmb_per_usd not in (None, 0)
                    else None
                ),
            }
        )
    wb.close()
    return pd.DataFrame(records)

pb_cvi = load_pb_cvi()
z2i = load_zero2ipo()
df = pb_cvi.merge(z2i, on="Year", how="outer").sort_values("Year")
df = df[(df["Year"] >= 2000) & (df["Year"] <= 2021)]
df.to_csv(INTERMEDIATE, index=False)

fig, ax = plt.subplots(figsize=(8, 5))
ax.plot(df["Year"], df["PitchBook"], color="#1f3b73", marker="o", lw=1.8, label="PitchBook")
ax.plot(
    df["Year"],
    df["CVSource / ChinaVenture"],
    color="#c0392b",
    marker="s",
    lw=1.8,
    label="CVSource / ChinaVenture",
)
ax.plot(df["Year"], df["Zero2IPO"], color="#e08a1e", marker="^", lw=1.8, label="Zero2IPO")
ax.set_xlabel("Year")
ax.set_ylabel("Chinese VC investment size (US$ billion)")
ax.set_xlim(2000, 2021)
ax.set_ylim(0, None)
ax.grid(True, alpha=0.3)
ax.legend(frameon=False, loc="upper left")
for spine in ("top", "right"):
    ax.spines[spine].set_visible(False)
fig.tight_layout()
fig.savefig(OUT)
print("wrote", OUT)
