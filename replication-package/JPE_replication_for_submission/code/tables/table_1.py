import csv
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "auxiliary"))
from paths import ANALYSIS_DIR, INTERMEDIATE_DATA_DIR, INTERMEDIATE_DIR, MANUAL_DIR, MANUAL_DIR_CONF, TABLE_DIR

PANEL_A_DTA = INTERMEDIATE_DATA_DIR / "pbdealinvestor_vc.dta"
PANEL_A_METADATA = MANUAL_DIR_CONF / "table1_panelA_vc_groups.csv"
PANEL_B_SOURCE = MANUAL_DIR / "table1_panelB_funding_sources.csv"
PANEL_C_WORKBOOK = MANUAL_DIR_CONF / "table1_panelC_Copy of VC Returns 2024Q4 v2.xlsx"
PANEL_C_REFERENCE = MANUAL_DIR_CONF / "table1_panelC_returns.csv"
PANEL_D_WORKBOOK = MANUAL_DIR_CONF / "table1_panelD_Chinese VC.xlsx"
PANEL_D_REFERENCE = MANUAL_DIR_CONF / "table1_panelD_billionaires.csv"
PANEL_B_GENERATED = INTERMEDIATE_DIR / "table1_panelB_funding_sources_generated.csv"
PANEL_C_GENERATED = INTERMEDIATE_DIR / "table1_panelC_returns_generated.csv"
PANEL_D_GENERATED = INTERMEDIATE_DIR / "table1_panelD_current_workbook_counts.csv"
PANEL_D_COMPARISON = INTERMEDIATE_DIR / "table1_panelD_published_vs_current_workbook.csv"
OUT = TABLE_DIR / "table1.csv"

PANEL_C_REGIONS = ["China", "India", "U.S.", "Europe", "Middle East", "Canada"]
PANEL_D_BINS = [
    ("Top Ten", 1, 10),
    ("Next Two Deciles (11--30)", 11, 30),
    ("Next Two Deciles (31--50)", 31, 50),
    ("Next Five Deciles (51--100)", 51, 100),
]

def _require_nonempty(path: Path) -> None:
    if not path.exists():
        raise FileNotFoundError(f"Required input not found: {path}")
    if path.stat().st_size == 0:
        raise ValueError(f"Required input is empty: {path}")

def _require_columns(df: pd.DataFrame, columns: list[str], source: Path) -> None:
    missing = [c for c in columns if c not in df.columns]
    if missing:
        raise ValueError(f"{source} is missing required columns: {missing}")

def _write_intermediate(df: pd.DataFrame, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(path, index=False)

def _assert_matches_reference(
    generated: pd.DataFrame,
    reference_path: Path,
    columns: list[str],
    label: str,
) -> None:
    _require_nonempty(reference_path)
    reference = pd.read_csv(reference_path)
    _require_columns(reference, columns, reference_path)
    left = generated[columns].reset_index(drop=True)
    right = reference[columns].reset_index(drop=True)
    try:
        pd.testing.assert_frame_equal(left, right, check_dtype=False, atol=1e-9, rtol=1e-9)
    except AssertionError as exc:
        raise ValueError(
            f"{label} generated values differ from the reference values in {reference_path}"
        ) from exc

def build_panel_a() -> pd.DataFrame:
    _require_nonempty(PANEL_A_DTA)
    metadata = pd.read_csv(PANEL_A_METADATA)

    usecols = ["investorid", "investorname", "hqcountry_inv", "dealsize", "year"]
    pieces = []
    for chunk in pd.read_stata(
        PANEL_A_DTA,
        columns=usecols,
        chunksize=250_000,
        convert_categoricals=False,
    ):
        chunk = chunk[
            chunk["year"].between(2000, 2019)
            & chunk["hqcountry_inv"].isin(["China", "Hong Kong"])
        ]
        if chunk.empty:
            continue
        pieces.append(
            chunk.groupby(
                ["investorid", "investorname", "hqcountry_inv"],
                as_index=False,
                dropna=False,
            ).agg(
                dealsize=("dealsize", "sum"),
                deals=("investorname", "size"),
            )
        )

    if not pieces:
        raise ValueError(f"No Panel A observations after upstream filters: {PANEL_A_DTA}")

    investors = (
        pd.concat(pieces, ignore_index=True)
        .groupby(["investorid", "investorname", "hqcountry_inv"], as_index=False)
        .agg(dealsize=("dealsize", "sum"), deals=("deals", "sum"))
    )
    upstream_top100 = investors.sort_values("dealsize", ascending=False).head(100)
    top5 = upstream_top100.sort_values(
        ["deals", "dealsize", "investorname"],
        ascending=[False, False, True],
    ).head(5)

    panel = top5.rename(columns={"investorname": "group"})[["group", "deals"]]
    panel = panel.merge(
        metadata.drop(columns=["deals"], errors="ignore"),
        on="group",
        how="left",
        validate="1:1",
    )
    if panel.isna().any(axis=None):
        missing = panel[panel.isna().any(axis=1)]["group"].tolist()
        raise ValueError(f"Missing Panel A classification metadata for: {missing}")

    if "deals" in metadata.columns:
        expected = metadata.set_index("group")["deals"].astype(int).to_dict()
        mismatches = [
            (row.group, int(row.deals), expected.get(row.group))
            for row in panel.itertuples(index=False)
            if expected.get(row.group) is not None and int(row.deals) != expected[row.group]
        ]
        if mismatches:
            raise ValueError(
                "Generated Panel A deal counts differ from the reference values: "
                f"{mismatches}"
            )

    return panel

def build_panel_b() -> pd.DataFrame:
    _require_nonempty(PANEL_B_SOURCE)
    panel = pd.read_csv(PANEL_B_SOURCE)
    columns = ["source", "received_any_pct", "share_of_funding_pct"]
    _require_columns(panel, columns, PANEL_B_SOURCE)
    panel = panel[columns].copy()
    panel["received_any_pct"] = pd.to_numeric(panel["received_any_pct"])
    panel["share_of_funding_pct"] = pd.to_numeric(panel["share_of_funding_pct"])

    expected_sources = ["Government VC", "Private VC"]
    if panel["source"].tolist() != expected_sources:
        raise ValueError(
            "Panel B source rows changed; expected Government VC then Private VC, "
            f"got {panel['source'].tolist()}"
        )

    _write_intermediate(panel, PANEL_B_GENERATED)
    return panel

def build_panel_c() -> pd.DataFrame:
    _require_nonempty(PANEL_C_WORKBOOK)
    wb = load_workbook(PANEL_C_WORKBOOK, read_only=True, data_only=True)
    if "Tables" not in wb.sheetnames:
        raise ValueError(f"Panel C workbook lacks a 'Tables' sheet: {PANEL_C_WORKBOOK}")
    rows = list(wb["Tables"].iter_rows(values_only=True))

    table_idx = None
    for i, row in enumerate(rows):
        first = row[0] if row else None
        if (
            isinstance(first, str)
            and first.startswith("Table 3:")
            and "2007q3 - 2019q4" in first
        ):
            table_idx = i
            break
    if table_idx is None:
        raise ValueError(f"Could not find State Street Table 3 in {PANEL_C_WORKBOOK}")

    header = [str(x).strip() if x is not None else "" for x in rows[table_idx + 1]]
    try:
        region_col = header.index("Region")
        irr_col = header.index("IRR")
        pme_col = header.index("KS PME")
    except ValueError as exc:
        raise ValueError("Panel C State Street Table 3 header is incomplete") from exc

    parsed = []
    for row in rows[table_idx + 2:]:
        region = row[region_col]
        if region is None or str(region).strip() == "":
            break
        region = str(region).strip()
        if region.startswith("Table "):
            break
        parsed.append(
            {
                "region": region,
                "irr_pct": round(float(row[irr_col]) * 100, 2),
                "ks_pme": round(float(row[pme_col]), 2),
            }
        )

    table = pd.DataFrame(parsed)
    missing = [r for r in PANEL_C_REGIONS if r not in table["region"].tolist()]
    if missing:
        raise ValueError(f"Panel C State Street table is missing regions: {missing}")

    panel = (
        table[table["region"].isin(PANEL_C_REGIONS)]
        .assign(region=lambda d: pd.Categorical(d["region"], PANEL_C_REGIONS, ordered=True))
        .sort_values("region")
    )
    panel["region"] = panel["region"].astype(str)
    panel = panel[["region", "irr_pct", "ks_pme"]].reset_index(drop=True)

    if PANEL_C_REFERENCE.exists():
        _assert_matches_reference(
            panel,
            PANEL_C_REFERENCE,
            ["region", "irr_pct", "ks_pme"],
            "Panel C",
        )
    _write_intermediate(panel, PANEL_C_GENERATED)
    return panel

def _read_panel_d_records() -> pd.DataFrame:
    _require_nonempty(PANEL_D_WORKBOOK)
    wb = load_workbook(PANEL_D_WORKBOOK, read_only=True, data_only=True)
    if "task 6" not in wb.sheetnames:
        raise ValueError(f"Panel D workbook lacks a 'task 6' sheet: {PANEL_D_WORKBOOK}")
    ws = wb["task 6"]

    header_row = None
    headers = {}
    for row_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=20, max_col=8, values_only=True), 1):
        normalized = [str(v).strip() if v is not None else "" for v in row]
        if "Rank" in normalized and "VC backed?" in normalized:
            header_row = row_idx
            headers = {v: i for i, v in enumerate(normalized) if v}
            break
    if header_row is None:
        raise ValueError(f"Could not find Panel D header row in {PANEL_D_WORKBOOK}")

    records = []
    for row in ws.iter_rows(min_row=header_row + 1, max_col=8, values_only=True):
        try:
            rank = int(float(row[headers["Rank"]]))
        except (TypeError, ValueError):
            continue
        if not 1 <= rank <= 100:
            continue
        vc_raw = row[headers["VC backed?"]]
        records.append(
            {
                "rank": rank,
                "company": row[headers.get("Company", 2)],
                "name": row[headers.get("Name", 3)],
                "vc_backed": str(vc_raw).strip().lower() == "yes",
            }
        )

    records_df = pd.DataFrame(records)
    if len(records_df) != 100:
        raise ValueError(
            f"Panel D expected 100 Hurun rank<=100 entries, found {len(records_df)}"
        )
    return records_df

def _panel_d_counts(records: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for group, lo, hi in PANEL_D_BINS:
        subset = records[records["rank"].between(lo, hi)]
        total = len(subset)
        if total == 0:
            raise ValueError(f"Panel D bin has no records: {group}")
        vc_count = int(subset["vc_backed"].sum())
        rows.append(
            {
                "group": group,
                "rank_min": lo,
                "rank_max": hi,
                "entries": total,
                "vc_backed_entries": vc_count,
                "share_pct": round(100 * vc_count / total, 1),
            }
        )
    return pd.DataFrame(rows)

def build_panel_d() -> pd.DataFrame:
    generated = _panel_d_counts(_read_panel_d_records())
    _write_intermediate(generated, PANEL_D_GENERATED)

    if PANEL_D_REFERENCE.exists():
        reference = pd.read_csv(PANEL_D_REFERENCE)
        columns = ["group", "share_pct"]
        _require_columns(reference, columns, PANEL_D_REFERENCE)
        reference = reference[columns].copy()

        comparison = reference.assign(_order=range(len(reference))).merge(
            generated[["group", "entries", "vc_backed_entries", "share_pct"]],
            on="group",
            how="outer",
            suffixes=("_published", "_current_workbook"),
            validate="1:1",
        )
        comparison["difference_pct_points"] = (
            comparison["share_pct_current_workbook"] - comparison["share_pct_published"]
        ).round(1)
        comparison = comparison.sort_values("_order").drop(columns=["_order"])
        _write_intermediate(comparison, PANEL_D_COMPARISON)

        mismatched = comparison[
            comparison["difference_pct_points"].abs().fillna(0) > 1e-9
        ]
        if not mismatched.empty:
            print(
                "warning: Panel D current Hurun workbook counts differ from the "
                f"published values; using workbook-derived values for table1.csv and "
                f"writing a comparison to {PANEL_D_COMPARISON}"
            )
    return generated[["group", "share_pct"]]

def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)

    A = build_panel_a()
    B = build_panel_b()
    C = build_panel_c()
    D = build_panel_d()

    with OUT.open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["Panel A: Five Most Active China/Hong Kong Venture Groups, 2000-2019"])
        w.writerow(["Venture group", "Number of deals", "US parent/co-manager",
                    "Founder w/ US education & work", "Founder w/ US VC experience"])
        for r in A.itertuples(index=False):
            w.writerow(list(r))
        w.writerow([])
        w.writerow(["Panel B: Funding Sources of Chinese Companies (Government and/or Private VC), 2010-2023"])
        w.writerow(["Source", "Received any? (%)", "Share of total funding (%)"])
        for r in B.itertuples(index=False):
            w.writerow(list(r))
        w.writerow([])
        w.writerow(["Panel C: Financial Performance of VC Funds by Region, 2007-2019"])
        w.writerow(["Region", "IRR (%)", "KS PME"])
        for r in C.itertuples(index=False):
            w.writerow(list(r))
        w.writerow([])
        w.writerow(["Panel D: Share of Chinese Billionaires from VC-Backed Firms"])
        w.writerow(["Decile group", "Share (%)"])
        for r in D.itertuples(index=False):
            w.writerow(list(r))
    print("wrote", OUT)

if __name__ == "__main__":
    main()
